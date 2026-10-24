from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml.ns import qn
from statistics import mean
import json
R=Path('research-output');src=R/'BARANGAY_LUNA_CHAPTERS_1_TO_5_UPDATED_DFD.docx';d=Document(src)
a=json.loads((R/'work/survey-audit.json').read_text(encoding='utf-8'));ratings=a['ratings']
res=[k for k in ratings if k.startswith('R')];staff=['TS1','TS2'];admin=['TA1'];allids=res+staff+admin

def vals(ids,dim=None):
 out=[]
 for k in ids:
  rows=ratings[k] if dim is None else [ratings[k][dim]]
  out.extend(v for row in rows for v in row if isinstance(v,(int,float)))
 return out
combined=mean(vals(allids));assert len(vals(allids))==388
changes={
50:f'A descriptive evaluation of the Barangay Luna Smart Portal included 13 questionnaires covering 10 residents, two barangay staff members, and one administrator. The participants completed their own survey questionnaires based on their experience using the system. The staff and administrator evaluation took place at the barangay hall. The resident assessment included 298 valid ratings across 30 items and obtained a pooled mean of 4.42, interpreted as agree. The staff and administrator assessment obtained a mean of 4.46 from 90 ratings. Across all groups, 388 valid ratings produced a combined mean of {combined:.2f}, interpreted as agree. Automated regression checks supported selected request, issuance, payment, template, and interface behavior. The findings support continued refinement and a controlled pilot. The study does not establish a measured reduction in processing time, security certification, or formal barangay acceptance.',
301:'The current evaluation included 10 resident questionnaires, two barangay staff questionnaires, and one administrator questionnaire. The actual participants completed their own questionnaire responses based on their experience using the system. Staff and administrator testing and questionnaire completion were conducted at the barangay hall. Results are reported by group and as a combined descriptive summary. The earlier April interviews remain historical requirements evidence and are distinct from the current system evaluation.',
309:'This chapter presents the implemented system, verified regression checks, and evaluation results from residents, barangay staff, and the administrator. The 13 questionnaires cover the three intended user groups. All participants completed their own questionnaires. The actual barangay staff and administrator used the system and filled out their questionnaires at the barangay hall. Development outcomes, test cases, user perceptions, operational readiness, implementation steps, and risks are discussed in relation to the current prototype.',
338:'Nine resident questionnaires list a cellphone and one lists a computer. The two staff questionnaires and the administrator questionnaire list a computer. Age and gender statistics are not reported for this evaluation because the supplied forms do not provide a complete demographic dataset. The earlier proposal’s gender counts belong to a separate preliminary survey and are not transferred to the current sample.',
339:'The resident dataset contains 300 possible item responses, of which 298 are valid numeric ratings and two are N/A. C1 and E5 therefore have nine resident ratings each; all other resident items have ten. The two staff questionnaires contain 60 valid ratings and the administrator questionnaire contains 30. All 13 questionnaires therefore provide 388 valid ratings out of 390 possible responses. The duplicate photograph was excluded before tallying. The item and evaluator audit tables are retained in Appendix E.',
363:'4.6.4 Staff Administrator and Combined Evaluation Results',
364:'Table 18  Evaluation results across all user groups',
366:f'The two barangay staff questionnaires obtained an overall mean of {mean(vals(staff)):.2f}, and the administrator questionnaire obtained {mean(vals(admin)):.2f}; both are interpreted as agree. Together, their 90 ratings produced a pooled mean of 4.46. The actual barangay staff and administrator personally completed these questionnaires while testing the system at the barangay hall. Their responses describe their assessment of the system’s functionality, usability, reliability, performance, observed security behavior, and overall satisfaction. Their results are presented by user group and included in the combined summary.',
372:f'The resident, staff, and administrator findings support a favorable initial assessment of the prototype among the participating users. Across the 13 questionnaires, the combined mean was {combined:.2f}, interpreted as agree. The actual staff and administrator completed their own questionnaires while using the system at the barangay hall. This initial evaluation is distinct from formal deployment acceptance; no signed acceptance form or official authorization for operational adoption is inferred from the questionnaire ratings.',
373:'A follow-up evaluation should record completion of concrete tasks: create and verify a resident account, find a profile, review a request, reject it with a reason, prepare a certificate, check parent and child wording, collect cash, confirm release, inspect reports, and restore a backup in a test environment. Record errors, assistance needed, and readiness decisions. Resident tests should also cover phone access, status comprehension, error correction, and the intended network conditions.',
387:f'The descriptive evaluation included 13 questionnaires from 10 residents, two barangay staff members, and one administrator. The actual participants personally completed their questionnaires based on their experience using the system. The staff and administrator evaluation was conducted at the barangay hall. The resident questionnaires provided 298 valid numeric ratings and two N/A responses, with an overall mean of 4.42. The staff mean was {mean(vals(staff)):.2f}, and the administrator mean was {mean(vals(admin)):.2f}. Across all groups, 388 valid ratings produced a combined mean of {combined:.2f}. These means are interpreted as agree. Selected functional and interface regression checks also passed in the current build. Formal deployment acceptance and production implementation remain separate from this initial evaluation.',
391:f'The barangay staff and administrator results indicate a favorable initial assessment of the office functions, with overall means of {mean(vals(staff)):.2f} and {mean(vals(admin)):.2f}, respectively. When all questionnaires are included, the combined mean of {combined:.2f} indicates agreement with the evaluation statements across the participating groups. These results reflect questionnaires personally completed by the actual staff and administrator during system testing at the barangay hall. The small sample limits how widely the findings can be generalized. The findings support further refinement and a controlled pilot, subject to the approved workflow, accurate data, training, connectivity, document conversion, and gateway configuration.',
395:'Review the evaluation findings and proposed improvements with the participating staff and administrator. Review resident registration rules, identity verification, document fees, template wording, approval responsibilities, and release procedures. Obtain a documented readiness decision and the required authorization before introducing live resident and transaction data into an operational deployment.',
402:'Expand the evaluation to additional residents and barangay personnel using clearly documented sampling and task procedures. Retain personally completed questionnaires and record task completion to support the analysis of questionnaire results. If operational improvement is studied, collect comparable baseline and portal measurements of processing time, follow-up visits, errors, and completion rates. Do not compare the preliminary feature-preference survey directly with the later usability questionnaire as though they were the same measure.',
578:'R01–R10 identify the handwritten resident numbers 1–10. TS1 and TS2 identify the two staff questionnaire sets, and TA1 identifies the administrator questionnaire set. These identifiers are retained for traceability to the source photographs. The actual staff and administrator personally completed their questionnaires while using the system at the barangay hall. The continuation pages are grouped with their adjoining cover pages in the supplied image sequence. Each row below lists the five item responses in item order. NA is excluded from numeric means. The duplicated photograph labelled Image 23 and Image 24 is counted only once.'
}
for i,value in changes.items():
 p=d.paragraphs[i]
 if p.runs:p.runs[0].text=value
 else:p.add_run(value)
 for run in p.runs[1:]:run.text=''
# Explain the combined computation directly after the group discussion.
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
node=OxmlElement('w:p');d.paragraphs[366]._p.addnext(node);p=Paragraph(node,d);p.style=d.paragraphs[366].style
p.add_run(f'The combined mean of {combined:.2f} is calculated from the sum of all 388 valid item ratings divided by 388. N/A responses are excluded. This calculation weights each valid item response equally rather than averaging rounded group means. Table 18 retains the separate group means to show the evaluation of each intended user group alongside the combined result. The small number of staff and administrator questionnaires limits generalization and does not establish differences between occupational groups.')
# Update composition while retaining the source ratings.
t=d.tables[9]
rows=[['Resident','10','Actual resident users; questionnaires completed by residents','Resident descriptive evaluation'],['Barangay staff','2','Actual barangay staff; personally completed questionnaires during system testing','Staff descriptive evaluation'],['Administrator','1','Actual administrator; personally completed questionnaire during system testing','Administrator descriptive evaluation'],['Total','13','10 residents, 2 staff members and 1 administrator','Group results and combined descriptive summary']]
for row,values in zip(t.rows[1:],rows):
 for c,v in zip(row.cells,values):c.text=v
# Add the resident and all-group columns to the existing results table.
t=d.tables[17];t.add_column(Inches(1.0))
widths=[1.65,.98,.98,.98,1.1]
for col,w in zip(t.columns,widths):col.width=Inches(w)
heads=['Dimension','Resident mean\nn = 10','Staff mean\nn = 2','Admin mean\nn = 1','Combined mean\nn = 13']
for c,v in zip(t.rows[0].cells,heads):c.text=v
for i,name in enumerate(a['dimension_names']):
 for c,v in zip(t.rows[i+1].cells,[name]+[f'{mean(vals(ids,i)):.2f}' for ids in [res,staff,admin,allids]]):c.text=v
for c,v in zip(t.rows[-1].cells,['Overall']+[f'{mean(vals(ids)):.2f}' for ids in [res,staff,admin,allids]]):c.text=v
for row in t.rows:
 for c,w in zip(row.cells,widths):
  c.width=Inches(w)
  for p in c.paragraphs:
   for run in p.runs:run.font.size=Pt(9)
for row in d.tables[18].rows:
 for c in row.cells:
  c.text=c.text.replace('TS1 team staff role','TS1 barangay staff').replace('TS2 team staff role','TS2 barangay staff').replace('TA1 team administrator role','TA1 administrator')
d.tables[19].cell(1,1).text='Obtain authorization and agree on a controlled pilot with the designated barangay personnel.'
d.tables[19].cell(4,0).text='Training and follow-up testing'
d.tables[8].cell(1,2).text='Replacement and modal regression checks passed; repeat with office users.'
# Refresh the table title without removing existing PAGEREF fields.
for p in d.paragraphs:
 for run in p.runs:
  if 'Table 18  Team role assessment reported separately' in run.text:run.text=run.text.replace('Table 18  Team role assessment reported separately','Table 18  Evaluation results across all user groups')
d.core_properties.comments='Personally completed questionnaires from residents, actual barangay staff, and the administrator; group and combined evaluation results.'
# Adviser-approved wording must survive unchanged and in order.
original=Document(r'C:/Users/Christian Deloy/Downloads/FINAL.docx')
approved=[p.text for p in original.paragraphs[118:429] if p.text.strip()];text=[p.text for p in d.paragraphs];cursor=0
for value in approved:cursor=text.index(value,cursor)+1
assert len(approved)==107
out=R/'BARANGAY_LUNA_CHAPTERS_1_TO_5_ALL_USER_EVALUATION.docx';d.save(out)
(R/'work/all-user-evaluation-audit.json').write_text(json.dumps({'output':str(out),'approved_paragraphs_preserved':len(approved),'valid_ratings':len(vals(allids)),'combined_mean':combined,'staff_mean':mean(vals(staff)),'administrator_mean':mean(vals(admin)),'recording_method':'Personally completed resident, staff, and administrator questionnaires; staff and administrator evaluation at the barangay hall, as clarified by the user'},indent=2),encoding='utf-8')
print('Saved',out,'Combined:',combined,'Approved paragraphs preserved:',len(approved))
