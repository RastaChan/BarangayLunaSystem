from pathlib import Path
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from shutil import copy2
import html,json
R=Path('research-output');W=R/'work';im=Image.new('RGB',(2600,1750),'white');dr=ImageDraw.Draw(im);svg=[]
def text(x,y,value,size=31,bold=False):
 f=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf' if bold else 'C:/Windows/Fonts/arial.ttf',size);ls=value.split('\n')
 for i,line in enumerate(ls):
  yy=y+(i-(len(ls)-1)/2)*(size+12);dr.text((x,yy),line,font=f,fill='#222',anchor='mm')
  svg.append(f'<text x="{x}" y="{yy}" font-family="Arial" font-size="{size}" font-weight="{700 if bold else 400}" text-anchor="middle" dominant-baseline="central">{html.escape(line)}</text>')
def box(r,rounded=False):
 if rounded:dr.rounded_rectangle(r,30,'#fff2a8','#333',3)
 else:dr.rectangle(r,'#fff2a8','#333',3)
 x,y,a,b=r;svg.append(f'<rect x="{x}" y="{y}" width="{a-x}" height="{b-y}" rx="{30 if rounded else 0}" fill="#fff2a8" stroke="#333" stroke-width="3"/>')
def line(points,arrow=True):
 dr.line(points,fill='#333',width=3)
 svg.append('<polyline points="'+' '.join(f'{x},{y}' for x,y in points)+'" fill="none" stroke="#333" stroke-width="3"'+(' marker-end="url(#a)"' if arrow else '')+'/>')
 if arrow:
  x,y=points[-1];a,b=points[-2];dx=x-a;dy=y-b
  if abs(dx)>abs(dy):s=1 if dx>0 else -1;tri=[(x,y),(x-s*18,y-9),(x-s*18,y+9)]
  else:s=1 if dy>0 else -1;tri=[(x,y),(x-9,y-s*18),(x+9,y-s*18)]
  dr.polygon(tri,fill='#333')
text(1300,35,'PROPOSED DATA CONTEXT DIAGRAM',43,True)
box((1100,100,1500,260));text(1300,180,'SMS Gateway',38,True)
box((70,710,450,950));text(260,830,'Resident',38,True)
box((2150,710,2530,950));text(2340,830,'Barangay Staff',38,True)
box((1100,1440,1500,1620));text(1300,1530,'Administrator',38,True)
box((950,650,1650,1020),True);line([(950,710),(1650,710)],False);text(1300,680,'0',34,True)
text(1300,865,'Smart Barangay Portal\nwith QR Profiling and SMS Alerts\nfor Barangay Luna',38,True)
line([(450,770),(950,770)]);text(690,620,'Login / profile details\nDocument request, purpose, copies\nComplaints and feedback',29)
line([(950,930),(450,930)]);text(690,1080,'Account access / QR profile\nRequest previews and status\nPosts and service updates',29)
line([(2150,770),(1650,770)]);text(1910,565,'Resident identity and requirements\nOnline / walk-in applicant details\nReview, preparation and print requests\nCash received / Confirm Release\nAnnouncements / complaint actions\nReport requests',29)
line([(1650,930),(2150,930)]);text(1910,1090,'Resident lookup and request queue\nPrintable documents / amounts due\nPayment and release records\nCommunity records and reports',29)
line([(1120,650),(1120,260)]);text(860,435,'Recipient mobile number\nand SMS message',29)
line([(1480,260),(1480,650)]);text(1740,330,'Gateway send response\nfor SMS logging',29)
line([(1120,1440),(1120,1020)]);text(865,1260,'Account and access changes\nTemplate / fee updates\nSettings and backup actions\nReport parameters',29)
line([(1480,1020),(1480,1440)]);text(1735,1260,'Administrative overview\nService and collection reports\nAudit records\nSettings / backup results',29)
text(1300,1695,'Single system process • External entities and data flows • Current implementation',28)
img=R/'PROPOSED_DATA_CONTEXT_DIAGRAM_CURRENT_SYSTEM.png';im.save(img)
head='<svg xmlns="http://www.w3.org/2000/svg" width="2600" height="1750" viewBox="0 0 2600 1750"><defs><marker id="a" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8" fill="#333"/></marker></defs><rect width="2600" height="1750" fill="white"/>'
(R/'PROPOSED_DATA_CONTEXT_DIAGRAM_CURRENT_SYSTEM.svg').write_text(head+''.join(svg)+'</svg>',encoding='utf-8')
p=R/'BARANGAY_LUNA_CHAPTERS_1_TO_5_ALL_USER_EVALUATION.docx';d=Document(W/'before-context-diagram-comment.docx');before=[x.text for x in d.paragraphs]
pa=next(x for x in d.paragraphs if x.text.startswith('The Proposed Data Context Diagram illustrates'))
v='The Proposed Data Context Diagram represents the Smart Barangay Portal with QR Profiling and SMS Alerts for Barangay Luna as a single system process interacting with the Resident, Barangay Staff, Administrator, and SMS Gateway. Residents provide login or profile details, document requests with their purpose and number of copies, complaints, and feedback. The system returns account access and a QR profile, request previews and status, published posts, and service updates. Barangay staff provide verified resident and applicant details, review and document preparation instructions, cash payment and release actions, announcements, complaint actions, and report requests. They receive resident lookup results, the request queue, printable documents, amounts due, payment and release records, community records, and reports. The administrator provides account and access changes, template and fee updates, settings and backup actions, and report parameters. The system returns administrative information, service and collection reports, audit records, and settings or backup results. The SMS Gateway receives recipient mobile numbers and messages and returns send responses for logging. Cash payments are recorded by staff at the office. The diagram shows the system boundary and its external data exchanges; the six internal processes and data stores are detailed in the Proposed Data Flow Diagram.'
pa.runs[0].text=v
for run in pa.runs[1:]:run.text=''
cap=next(x for x in d.paragraphs if x.text=='Figure 6  Original proposed context diagram');node=cap._p.getprevious()
for child in list(node):
 if child.tag!=qn('w:pPr'):node.remove(child)
pi=Paragraph(node,d);pi.add_run().add_picture(str(img),width=Inches(9.25));pi.paragraph_format.keep_with_next=True
from copy import deepcopy
from docx.oxml import OxmlElement
oldsect=next(x for x in d.paragraphs if x._p.find('.//'+qn('w:sectPr')) is not None and x.text.startswith('Process 6.0, Reports'))
prior=deepcopy(oldsect._p.find('.//'+qn('w:sectPr')))
pa._p.get_or_add_pPr().append(prior)
for x in oldsect._p.findall('.//'+qn('w:pgNumType')):x.attrib.pop(qn('w:start'),None)
land=deepcopy(d.element.body.sectPr)
pg=land.find(qn('w:pgSz'));pg.set(qn('w:w'),'16838');pg.set(qn('w:h'),'11906');pg.set(qn('w:orient'),'landscape')
mg=land.find(qn('w:pgMar'))
for k in ['top','bottom','left','right']:mg.set(qn('w:'+k),'720')
typ=land.find(qn('w:type'))
if typ is None:typ=OxmlElement('w:type');land.insert(0,typ)
typ.set(qn('w:val'),'nextPage')
cap._p.get_or_add_pPr().append(land)
cap.paragraph_format.keep_with_next=False
newcap='Figure 6  Proposed context diagram based on the implemented system'
for para in d.paragraphs:
 for run in para.runs:
  if 'Figure 6  Original proposed context diagram' in run.text:run.text=run.text.replace('Figure 6  Original proposed context diagram',newcap)
for i,x in enumerate(d.element.findall('.//'+qn('wp:docPr')),1):x.set('id',str(i))
d.save(p)
print('Updated context diagram, its explanation, caption, and figure-list entry.')
